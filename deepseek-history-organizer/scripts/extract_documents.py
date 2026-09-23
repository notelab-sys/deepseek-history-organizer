#!/usr/bin/env python3
"""Extract readable text from Word (.docx), PDF, Excel (.xlsx) and plain-text files.

Word (.docx) is the primary target. Extracted text is saved as JSON for use by
build_html.py (--documents) and optionally as readable .txt copies (--txt-dir).

Usage:
  python extract_documents.py <folder-or-file> [...] [-o docs_extracted.json] [--txt-dir docs]
"""

import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path


def find_officecli():
    exe = shutil.which("officecli")
    if exe:
        return exe
    candidates = [
        Path.home() / "AppData/Local/OfficeCLI/officecli.exe",
        Path("/usr/local/bin/officecli"),
    ]
    for c in candidates:
        if c.exists():
            return str(c)
    return None


def officecli_text(path):
    """Extract text via officecli view <file> text; return None on failure."""
    exe = find_officecli()
    if not exe:
        return None
    try:
        proc = subprocess.run(
            [exe, "view", str(path), "text", "--max-lines", "200000"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=180,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    if proc.returncode != 0:
        return None
    lines = []
    for raw in proc.stdout.splitlines():
        line = re.sub(r"^\[.*?\]\]\s*", "", raw).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def extract_docx_tables(path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table

    doc = Document(path)
    parts = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            table = Table(child, doc)
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    parts.append(line)
    return "\n".join(parts)


def extract_docx_python(path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    parts = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    parts.append(line)
    return "\n".join(parts)


def extract_docx(path):
    oc_text = officecli_text(path)
    if oc_text is None:
        return extract_docx_python(path)
    table_marker = re.compile(r"^\[Table:\s*\d+\s*rows?\]$")
    lines = oc_text.splitlines()
    if not any(table_marker.match(line) for line in lines):
        return oc_text
    lines = [line for line in lines if not table_marker.match(line)]
    table_text = extract_docx_tables(path)
    if table_text:
        lines.append("")
        lines.append("## 表格内容")
        lines.append(table_text)
    return "\n".join(lines)


def extract_pdf(path):
    text_parts = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                text_parts.append(text)
    except Exception:  # noqa: BLE001 - fall back to pdfplumber
        text_parts = []
    if text_parts:
        return "\n".join(text_parts)
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                text_parts.append(text)
    return "\n".join(text_parts)


def extract_xlsx_openpyxl(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            parts.append(f"## 工作表：{ws.title}")
            for row in ws.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None and str(v).strip()]
                if values:
                    parts.append(" | ".join(values))
    finally:
        wb.close()
    return "\n".join(parts)


def extract_xlsx(path):
    oc_text = officecli_text(path)
    if oc_text is not None:
        lines = []
        for row in oc_text.splitlines():
            values = []
            for cell in row.split("\t"):
                match = re.match(r"^[A-Z]+\d+=(.*)$", cell)
                values.append(match.group(1) if match else cell)
            cleaned = [v.strip() for v in values if v and v.strip()]
            if cleaned:
                lines.append(" | ".join(cleaned))
        if lines:
            return "\n".join(lines)
    return extract_xlsx_openpyxl(path)


def extract_pptx(path):
    oc_text = officecli_text(path)
    return oc_text or ""


def extract_plain(path):
    for encoding in ("utf-8-sig", "gbk", "latin-1"):
        try:
            return Path(path).read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return ""


EXTRACTORS = {
    ".docx": extract_docx,
    ".pdf": extract_pdf,
    ".xlsx": extract_xlsx,
    ".pptx": extract_pptx,
    ".txt": extract_plain,
    ".md": extract_plain,
    ".csv": extract_plain,
}


def collect_files(paths):
    files = []
    for item in paths:
        p = Path(item)
        if p.is_dir():
            for f in sorted(p.rglob("*")):
                if f.suffix.lower() in EXTRACTORS and f.is_file():
                    files.append(f)
        elif p.is_file() and p.suffix.lower() in EXTRACTORS:
            files.append(p)
        else:
            print(f"跳过不支持的文件: {p}")
    seen = set()
    result = []
    for f in files:
        key = str(f.resolve())
        if key not in seen:
            seen.add(key)
            result.append(f)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("inputs", nargs="+", help="文档文件或文件夹路径")
    parser.add_argument("-o", "--output", default="docs_extracted.json", help="输出 JSON 路径（默认 docs_extracted.json）")
    parser.add_argument("--txt-dir", default=None, help="可选：同时输出可读 .txt 副本到此目录")
    args = parser.parse_args()

    files = collect_files(args.inputs)
    if not files:
        sys.exit("没有找到可读取的文档（支持: " + ", ".join(sorted(EXTRACTORS)) + "）")

    docs = []
    skipped = []
    for f in files:
        try:
            text = EXTRACTORS[f.suffix.lower()](f).strip()
        except Exception as exc:  # noqa: BLE001 - report and continue
            skipped.append(f"{f.name}: {exc}")
            continue
        if not text:
            skipped.append(f"{f.name}: 未提取到文字")
            continue
        docs.append({"name": f.name, "path": str(f), "text": text, "char_count": len(text)})

    if not docs:
        detail = "\n".join(skipped) if skipped else ""
        sys.exit("未能从任何文档中提取到文字" + (f"\n{detail}" if detail else ""))

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(docs, ensure_ascii=False, indent=2), encoding="utf-8")

    if args.txt_dir:
        txt_dir = Path(args.txt_dir)
        txt_dir.mkdir(parents=True, exist_ok=True)
        for d in docs:
            (txt_dir / f"{Path(d['name']).stem}.txt").write_text(d["text"], encoding="utf-8")

    print(f"成功提取: {len(docs)} 个文档")
    for d in docs:
        print(f"  - {d['name']}（{d['char_count']} 字）")
    if skipped:
        print(f"跳过 {len(skipped)} 个:")
        for s in skipped:
            print(f"  - {s}")
    print(f"已保存: {out_path}")
    if args.txt_dir:
        print(f"文本副本: {args.txt_dir}")


if __name__ == "__main__":
    main()
